import sys,copy,docx
from docx.oxml.ns import qn
src,out=sys.argv[1],sys.argv[2]
d=docx.Document(src);P=d.paragraphs
def setp(i,parts,font=None):
    """parts: list of (text,bold). Reuse first run's formatting."""
    p=P[i];base=p.runs[0]
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    base.text=parts[0][0]
    if parts[0][1] is not None: base.bold=parts[0][1]
    if font: base.font.name=font
    for t,b in parts[1:]:
        nr=copy.deepcopy(base._r);p._p.append(nr)
        r=p.runs[-1];r.text=t;r.bold=b
setp(2,[("Quick-reference visual identity guide  ·  For internal & agency use  ·  v2 (Oct 2026, matched to wrecktheroom.com)",None)])
setp(4,[("Wreck the Room: ",True),("the core emotional transformation brand. Black and warm white, with gold as the frame and red for emphasis. Used for wrecktheroom.com, keynote, coaching, mission content and hero brand moments.",False)])
setp(5,[("Architects: ",True),("the identity/community brand for people doing the work. Black, white and blue. Used for the community, the DIY program and program pages like /rebuild. (Naming note: the community is now just \"Architects,\" not \"Activated Architects.\")",False)])
setp(7,[("On Wreck the Room assets, gold is the frame (buttons, labels, lines, borders) and red is the punch (one or two highlighted words per headline). Do not mix red and blue as co-equal accents in the same asset. They mark different brand layers (Wreck vs. Build).",None)])
setp(9,[("Anton",True)],font="Anton")
setp(11,[("Used for all hero headlines, titles and bold statement copy, set in uppercase. This is the headline face live on wrecktheroom.com and /rebuild.",None)])
setp(12,[("Manrope",True)],font="Manrope")
setp(13,[("Body copy, captions, subheads and UI text. Labels and eyebrows: small uppercase Manrope with wide letter-spacing, in gold.",None)])
P[14].runs[0].font.name="Manrope"
setp(17,[("Architects hero lockup: white/blue mark and wordmark on black, concrete texture and blueprint linework, with aspirational interior or nature imagery. Never place the mark on light backgrounds.",None)])
setp(19,[("Wreck the Room hero lockup: B&W photography, cracked-concrete texture, distressed red wordmark, Jody's signature.",None)])
setp(20,[("VOICE & TONE: QUICK RULES",None)])

rows=[("0A0908","Deep Black","#0A0908","Core background, every brand"),
("F3EDE2","Warm White","#F3EDE2","Headlines and key text"),
("D2CBBE","Body Text","#D2CBBE","Paragraph copy on dark backgrounds"),
("C9A84C","Brand Gold","#C9A84C","Wreck the Room primary accent: buttons, labels, dividers, frames. Buttons use the gold gradient #ECD69B to #C9A84C to #A5862F"),
("ECD69B","Gold Highlight","#ECD69B","Light end of the gold gradient, hover states, small gold text"),
("C21F1F","Destruction Red","#C21F1F","Wreck the Room secondary accent: borders, tints, darker red moments"),
("D72626","Bright Signal Red","#D72626","Wreck the Room highlight: key words in headlines, emphasis"),
("4BA7F9","Electric Blue","#4BA7F9","Architects only: community, DIY program, /rebuild CTAs")]
t=d.tables[0];tbl=t._tbl;tmpl=copy.deepcopy(t.rows[1]._tr)
for r in list(t.rows): tbl.remove(r._tr)
for fill,name,hexv,use in rows:
    tr=copy.deepcopy(tmpl);tbl.append(tr)
    row=t.rows[-1]
    shd=row.cells[0]._tc.tcPr.find(qn('w:shd'))
    if shd is not None: shd.set(qn('w:fill'),fill)
    for ci,txt in ((1,name),(2,hexv),(3,use)):
        p=row.cells[ci].paragraphs[0];rs=p.runs
        for r in rs[1:]: r._r.getparent().remove(r._r)
        rs[0].text=txt
d.save(out)
txt='\n'.join(p.text for p in docx.Document(out).paragraphs)+'\n'.join(c.text for r in docx.Document(out).tables[0].rows for c in r.cells)
print('em dashes:',txt.count('—'))
